from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
OUTPUT_FILE = "Social_Media_Test_Cases.xlsx"
HEADERS = ["Test Case ID", "Module", "Test Scenario", "Test Steps", "Test Data", "Expected Result", "Actual Result", "Status", "Priority", "Remarks"]
WIDTHS = [14, 18, 38, 50, 32, 45, 28, 15, 11, 22]
CASES = [
    ("Registration", "Register with valid details", "1. Open sign-up page\n2. Enter name, email, password\n3. Click Sign Up", "Name: Test User\nEmail: test@example.com\nPassword: Test@1234", "Account is created, verification email is sent and user lands on onboarding/home page", "High"),
    ("Registration", "Register with an already registered email", "1. Open sign-up page\n2. Enter an existing email\n3. Click Sign Up", "Email: existing@example.com", "Error message 'Email already registered' is displayed", "High"),
    ("Registration", "Register with empty mandatory fields", "1. Open sign-up page\n2. Leave all fields blank\n3. Click Sign Up", "All fields empty", "Validation messages appear for every mandatory field", "High"),
    ("Registration", "Register with invalid email format", "1. Enter an invalid email\n2. Fill other fields\n3. Click Sign Up", "Email: test@@abc", "Error message 'Enter a valid email address' is displayed", "Medium"),
    ("Registration", "Register with a weak password", "1. Enter a short or simple password\n2. Click Sign Up", "Password: 123", "Password strength error is shown and account is not created", "High"),
    ("Registration", "Confirm password does not match", "1. Enter password\n2. Enter a different confirm password\n3. Click Sign Up", "Password: Test@1234\nConfirm: Test@4321", "Error message 'Passwords do not match' is displayed", "Medium"),
    ("Login", "Login with valid credentials", "1. Open login page\n2. Enter username and password\n3. Click Login", "Valid registered user", "User is logged in and the home feed is displayed", "High"),
    ("Login", "Login with incorrect password", "1. Enter valid username\n2. Enter wrong password\n3. Click Login", "Password: Wrong@123", "Error message 'Incorrect password' is displayed and login is denied", "High"),
    ("Login", "Login with unregistered username", "1. Enter unregistered username\n2. Enter any password\n3. Click Login", "User: nouser999", "Error message 'Account not found' is displayed", "Medium"),
    ("Login", "Login with empty username and password", "1. Leave both fields blank\n2. Click Login", "Empty fields", "Validation messages are shown for both fields", "Medium"),
    ("Login", "Forgot password flow", "1. Click 'Forgot password'\n2. Enter registered email\n3. Open reset link from email\n4. Set new password", "Registered email", "Reset email is received and user can log in with the new password", "High"),
    ("Login", "Account lock after repeated failed attempts", "1. Enter wrong password 5 times in a row", "Valid username, wrong password", "Account is temporarily locked and a lockout message is displayed", "Medium"),
    ("Login", "Login using Google/Facebook account", "1. Click 'Continue with Google'\n2. Select account and authorize", "Valid Google account", "User is logged in and linked to the social account", "Medium"),
    ("Login", "Session persists with Remember Me", "1. Login with Remember Me checked\n2. Close and reopen browser", "Valid user", "User remains logged in", "Low"),
    ("Profile", "Update profile picture with valid image", "1. Open profile\n2. Click Edit\n3. Upload JPG image under 5 MB\n4. Save", "profile.jpg (2 MB)", "Profile picture is updated and visible across the site", "Medium"),
    ("Profile", "Upload unsupported file as profile picture", "1. Open profile\n2. Upload an .exe file as picture", "setup.exe", "Error message about unsupported file type is displayed", "Medium"),
    ("Profile", "Edit name and bio", "1. Open Edit Profile\n2. Change name and bio\n3. Save", "Name: New Name\nBio: Hello world", "Updated details are saved and shown on the profile", "Medium"),
    ("Profile", "Change password with correct old password", "1. Open Settings > Security\n2. Enter old and new password\n3. Save", "Old: Test@1234\nNew: Test@5678", "Password is changed and user can log in with the new password", "High"),
    ("Post", "Create a text post", "1. Click Create Post\n2. Type text\n3. Click Post", "Text: Hello everyone", "Post is published and appears at the top of the feed", "High"),
    ("Post", "Create a post with an image", "1. Click Create Post\n2. Attach an image\n3. Add caption\n4. Post", "photo.png (1 MB)", "Post with image is published successfully", "High"),
    ("Post", "Create a post with a video", "1. Click Create Post\n2. Attach a supported video\n3. Post", "video.mp4 (20 MB)", "Video uploads, processes and plays correctly in the feed", "Medium"),
    ("Post", "Publish an empty post", "1. Click Create Post\n2. Leave content blank\n3. Click Post", "No content", "Post button is disabled or an error message is displayed", "Medium"),
    ("Post", "Post exceeding the character limit", "1. Click Create Post\n2. Paste text longer than the allowed limit\n3. Post", "Text of 5000+ characters", "Character limit message is shown and the post is blocked", "Medium"),
    ("Post", "Edit own post", "1. Open own post menu\n2. Click Edit\n3. Change text\n4. Save", "Updated text", "Post is updated and marked as edited", "Medium"),
    ("Post", "Delete own post", "1. Open own post menu\n2. Click Delete\n3. Confirm", "Existing post", "Post is removed from profile and feed", "High"),
    ("Post", "Hashtag and mention links", "1. Create a post with #tag and @username\n2. Click the hashtag and mention", "Text: Hi @user #testing", "Hashtag opens tag results and mention opens the user's profile", "Low"),
    ("Feed", "Home feed loads latest posts", "1. Login\n2. Open home page", "Following at least 3 users", "Latest posts from followed users are displayed in order", "High"),
    ("Feed", "Infinite scroll / pagination", "1. Open feed\n2. Scroll to the bottom", "Feed with 50+ posts", "More posts load without duplicates or errors", "Medium"),
    ("Interaction", "Like a post", "1. Open a post\n2. Click Like", "Any post", "Like count increases by 1 and the icon changes state", "High"),
    ("Interaction", "Unlike a post", "1. Click Like on a liked post again", "Previously liked post", "Like count decreases by 1", "Medium"),
    ("Interaction", "Comment on a post", "1. Open a post\n2. Type a comment\n3. Click Send", "Comment: Nice post!", "Comment is displayed under the post with the user's name", "High"),
    ("Interaction", "Delete own comment", "1. Open own comment menu\n2. Click Delete", "Existing comment", "Comment is removed and count updates", "Medium"),
    ("Interaction", "Share a post", "1. Click Share on a post\n2. Choose Share to feed", "Any public post", "Post is shared and appears on the user's profile", "Medium"),
    ("Follow", "Follow a user", "1. Open another user's profile\n2. Click Follow", "Public user account", "Button changes to Following and user's posts appear in feed", "High"),
    ("Follow", "Unfollow a user", "1. Open followed user's profile\n2. Click Following > Unfollow", "Followed user", "User is unfollowed and posts no longer appear in feed", "Medium"),
    ("Follow", "Send follow request to a private account", "1. Open a private profile\n2. Click Follow", "Private user account", "Request is sent and button shows Requested", "Medium"),
    ("Follow", "Block a user", "1. Open user's profile\n2. Click Block\n3. Confirm", "Any user", "Blocked user cannot view profile, message or follow", "High"),
    ("Messaging", "Send a text message", "1. Open Messages\n2. Select a friend\n3. Type and send", "Message: Hi", "Message is delivered and shown in the conversation", "High"),
    ("Messaging", "Send a message with attachment", "1. Open a chat\n2. Attach a file\n3. Send", "document.pdf (2 MB)", "Attachment is delivered and can be opened by the receiver", "Medium"),
    ("Messaging", "Real-time delivery and read status", "1. Send message from User A\n2. Observe User B", "Two logged-in users", "Message appears instantly for User B and the read status updates", "Medium"),
    ("Notifications", "Notification on like and comment", "1. User B likes and comments on User A's post\n2. Check User A notifications", "Two users", "User A receives notifications for the like and the comment", "Medium"),
    ("Notifications", "Mark notifications as read", "1. Open notifications\n2. Click Mark all as read", "Unread notifications", "Unread badge count becomes zero", "Low"),
    ("Search", "Search for an existing user", "1. Click search box\n2. Type a username", "Username: testuser", "Matching user appears in the results", "High"),
    ("Search", "Search for a non-existent user", "1. Type a random string in search", "Text: zzzxxyy123", "'No results found' message is displayed", "Low"),
    ("Search", "Search by hashtag", "1. Search for #testing", "Hashtag: #testing", "Posts using the hashtag are listed", "Medium"),
    ("Privacy", "Make profile private", "1. Open Settings > Privacy\n2. Enable Private Account", "Public account", "Only approved followers can see posts", "Medium"),
    ("Privacy", "Report a post", "1. Open post menu\n2. Click Report\n3. Select reason and submit", "Any post", "Confirmation message is shown and the report is logged", "Medium"),
    ("Privacy", "Deactivate account", "1. Open Settings > Account\n2. Click Deactivate\n3. Confirm with password", "Valid password", "Account is deactivated and profile is hidden", "Medium"),
    ("Logout", "Logout from the application", "1. Click profile menu\n2. Click Logout", "Logged-in user", "User is logged out and redirected to the login page", "High"),
    ("Logout", "Browser Back button after logout", "1. Logout\n2. Click the browser Back button", "Logged-out session", "Protected pages are not accessible and login page is shown", "High"),
    ("Security", "SQL injection in login field", "1. Enter SQL payload in username\n2. Click Login", "' OR '1'='1", "Input is rejected and no unauthorized access is granted", "High"),
    ("Security", "XSS script in post content", "1. Create a post with a script tag\n2. View the post", "<script>alert('xss')</script>", "Script is escaped and not executed", "High"),
    ("Security", "Session timeout on inactivity", "1. Login\n2. Stay idle beyond timeout\n3. Perform an action", "Idle for 30 minutes", "Session expires and user is asked to login again", "Medium"),
    ("Performance", "Home page load time", "1. Open the home page on a normal connection", "Broadband connection", "Page loads within 3 seconds", "Medium"),
    ("Compatibility", "Responsive layout on mobile viewport", "1. Open the site at 375 px width\n2. Navigate main pages", "Mobile viewport", "Layout adapts with no overlap or horizontal scroll", "Medium"),
    ("Compatibility", "Cross-browser check", "1. Open the site in Edge, Brave and Firefox\n2. Run login and post flow", "Edge, Brave, Firefox", "Features work consistently in all browsers", "Medium"),
]
thin = Side(style="thin", color="999999")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook()
ws = wb.active
ws.title = "Test Cases"
ws.append(HEADERS)
for index, (module, scenario, steps, data, expected, priority) in enumerate(CASES, start=1):
    ws.append([f"TC{index:03d}", module, scenario, steps, data, expected, "", "Not Executed", priority, ""])
for col, width in enumerate(WIDTHS, start=1):
    ws.column_dimensions[get_column_letter(col)].width = width
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = border
ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
status_validation = DataValidation(type="list", formula1='"Not Executed,Pass,Fail,Blocked"', allow_blank=False)
priority_validation = DataValidation(type="list", formula1='"High,Medium,Low"', allow_blank=False)
ws.add_data_validation(status_validation)
ws.add_data_validation(priority_validation)
last_row = ws.max_row
status_validation.add(f"H2:H{last_row}")
priority_validation.add(f"I2:I{last_row}")
summary = wb.create_sheet("Summary")
summary.append(["Metric", "Count"])
summary.append(["Total Test Cases", f"=COUNTA('Test Cases'!A2:A{last_row})"])
for label in ["Pass", "Fail", "Blocked", "Not Executed"]:
    summary.append([label, f"=COUNTIF('Test Cases'!H2:H{last_row},\"{label}\")"])
summary.append(["Pass %", "=IF(B2=0,0,B3/B2)"])
summary["B7"].number_format = "0.0%"
summary.column_dimensions["A"].width = 22
summary.column_dimensions["B"].width = 14
for cell in summary[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
for row in summary.iter_rows(min_row=1, max_row=summary.max_row):
    for cell in row:
        cell.border = border
wb.save(OUTPUT_FILE)
print(f"Created {OUTPUT_FILE} with {len(CASES)} test cases")